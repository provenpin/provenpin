from docx import Document
from report import build_report, has_page_break, headings, iter_block_items

COLUMNS = ["name", "qty"]


def make(tmp_path):
    path = tmp_path / "report.docx"
    build_report(path, "Q3 Report", COLUMNS, [["apples", 3], ["pears", 5]])
    return Document(str(path))


def test_title_and_headings(tmp_path):
    doc = make(tmp_path)
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.paragraphs[0].text == "Q3 Report"
    assert headings(doc) == [("Heading 1", "Summary"), ("Heading 1", "Appendix")]


def test_table_style_rows_and_cells(tmp_path):
    table = make(tmp_path).tables[0]
    assert table.style.name == "Light Grid Accent 1"
    assert len(table.rows) == 3
    assert [c.text for c in table.rows[0].cells] == COLUMNS
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True


def test_values_are_stringified(tmp_path):
    table = make(tmp_path).tables[0]
    assert table.cell(2, 1).text == "5"


def test_body_order_interleaves_tables_and_paragraphs(tmp_path):
    kinds = [type(b).__name__ for b in iter_block_items(make(tmp_path))]
    assert kinds == ["Paragraph", "Paragraph", "Paragraph", "Table", "Paragraph", "Paragraph", "Paragraph"]
    tables = [b for b in iter_block_items(make(tmp_path)) if type(b).__name__ == "Table"]
    assert len(tables) == 1


def test_page_break_is_an_empty_paragraph(tmp_path):
    doc = make(tmp_path)
    broken = [p for p in doc.paragraphs if has_page_break(p)]
    assert len(broken) == 1
    assert broken[0].text == ""


def test_unknown_table_style_raises_keyerror(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    try:
        table.style = "No Such Style"
    except KeyError:
        pass
    else:
        raise AssertionError("expected KeyError")
