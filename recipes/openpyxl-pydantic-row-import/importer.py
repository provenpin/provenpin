from datetime import date
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError


class Product(BaseModel):
    model_config = ConfigDict(extra="ignore", str_strip_whitespace=True)

    sku: str
    name: str
    quantity: int = Field(ge=0)
    price: float
    added: date | None = None


def rows_to_products(worksheet) -> list[Product]:
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {cell: i for i, cell in enumerate(header) if isinstance(cell, str)}
    products = []
    for row_number, row in enumerate(rows, start=2):
        raw = {name: row[i] if i < len(row) else None for name, i in columns.items()}
        try:
            products.append(Product.model_validate(raw))
        except ValidationError as exc:
            raise ValueError(f"row {row_number}: {exc}") from exc
    return products


def load_products(path: Any) -> list[Product]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return rows_to_products(workbook.active)
    finally:
        workbook.close()
