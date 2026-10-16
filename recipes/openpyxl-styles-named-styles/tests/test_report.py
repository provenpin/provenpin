import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from report import build_report, dumps, header_style, loads


def test_header_named_style_applied():
    ws = build_report([("widget", 1234.5)])["Report"]
    top = ws["A1"]
    assert top.style == "Header"
    assert top.font.bold
    assert top.font.color.rgb == "00FFFFFF"
    assert top.fill.fill_type == "solid"
    assert top.fill.fgColor.rgb == "004472C4"


def test_money_style_and_freeze_panes_survive_roundtrip():
    ws = loads(dumps(build_report([("widget", 1234.5)])))["Report"]
    assert ws.freeze_panes == "A2"
    amount = ws["B2"]
    assert amount.style == "Money"
    assert amount.number_format == "#,##0.00"
    assert amount.value == 1234.5


def test_unregistered_style_name_raises():
    ws = Workbook().active
    with pytest.raises(ValueError):
        ws["A1"].style = "Header"


def test_duplicate_style_name_rejected():
    wb = Workbook()
    wb.add_named_style(header_style())
    with pytest.raises(ValueError):
        wb.add_named_style(header_style())


def test_freeze_panes_a1_means_no_freeze():
    ws = Workbook().active
    ws.freeze_panes = "B2"
    assert ws.freeze_panes == "B2"
    ws.freeze_panes = "A1"
    assert ws.freeze_panes is None


def test_mutating_style_does_not_restyle_cells():
    wb = Workbook()
    style = header_style()
    wb.add_named_style(style)
    ws = wb.active
    ws["A1"] = "x"
    ws["A1"].style = "Header"
    style.font = Font(bold=True, italic=True)
    assert ws["A1"].font.italic is not True
