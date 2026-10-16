from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill


def header_style() -> NamedStyle:
    return NamedStyle(
        name="Header",
        font=Font(bold=True, color="FFFFFF"),
        fill=PatternFill("solid", fgColor="4472C4"),
        alignment=Alignment(horizontal="center"),
    )


def money_style() -> NamedStyle:
    return NamedStyle(name="Money", number_format="#,##0.00")


def build_report(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    wb.add_named_style(header_style())
    wb.add_named_style(money_style())
    ws.append(["item", "amount"])
    for cell in ws[1]:
        cell.style = "Header"
    for item, amount in rows:
        ws.append([item, amount])
        ws.cell(row=ws.max_row, column=2).style = "Money"
    ws.freeze_panes = "A2"
    return wb


def dumps(wb) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def loads(data) -> Workbook:
    return load_workbook(BytesIO(data))
