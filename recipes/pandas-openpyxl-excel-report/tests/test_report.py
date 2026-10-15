import pandas as pd
import pytest
from openpyxl import load_workbook

from report import read_report, write_report


@pytest.fixture
def df():
    return pd.DataFrame({
        "name": ["widget", "gadget"],
        "qty": [3, 10],
        "amount": [1234.5, None],  # None becomes an empty cell, read back as NaN
        "when": [pd.Timestamp("2024-01-02"), pd.Timestamp("2024-02-03")],
    })


def test_roundtrip_values_and_dtypes(tmp_path, df):
    path = write_report(df, tmp_path / "r.xlsx", money_cols=("amount",))
    back = read_report(path)
    pd.testing.assert_frame_equal(back, df, check_dtype=False)
    assert pd.api.types.is_datetime64_any_dtype(back["when"])


def test_styles_applied(tmp_path, df):
    path = write_report(df, tmp_path / "r.xlsx", money_cols=("amount",))
    ws = load_workbook(path)["Report"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.start_color.rgb.endswith("DDEBF7")
    assert ws["C2"].number_format == "#,##0.00"
    assert ws["B2"].number_format == "General"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D3"


def test_empty_cells_and_nan(tmp_path, df):
    path = write_report(df, tmp_path / "r.xlsx")
    ws = load_workbook(path)["Report"]
    assert ws["C3"].value is None
    assert pd.isna(read_report(path)["amount"].iloc[1])


def test_invalid_sheet_name_raises(tmp_path, df):
    with pytest.raises(ValueError):
        write_report(df, tmp_path / "r.xlsx", sheet="Q1/Totals")
