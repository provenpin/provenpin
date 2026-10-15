from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="DDEBF7")
MONEY_FORMAT = "#,##0.00"


def write_report(df: pd.DataFrame, path, sheet: str = "Report", money_cols=()) -> Path:
    path = Path(path)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=sheet, index=False)
        ws = xw.sheets[sheet]  # styles applied before the file is saved on exit

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

        for idx, col in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
            ws.column_dimensions[letter].width = longest + 2
            if col in money_cols:
                for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    cell.number_format = MONEY_FORMAT

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    return path


def read_report(path, sheet: str = "Report") -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
